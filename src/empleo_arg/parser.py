from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path
from typing import Iterable

from .model import Observation, QualityRecord, SourceSnapshot, decimal_or_none
from .periods import canonical_period
from .registry import CORE_INDICATORS, match_geography, match_indicator, norm


class SchemaDriftError(RuntimeError):
    pass


_LAYOUT_PERIOD_STYLES = {
    "indec_historical_matrix_v1": ("compact_roman", "date", "canonical"),
    "indec_current_matrix_v2": ("quarter_text", "canonical", "date", "compact_roman"),
    "indec_quality_v1": ("quarter_text", "canonical", "date", "compact_roman"),
}

_MISSING_MARKERS = {"s d", "sd", "s/d", "...", "..", "-", "///", "sin dato", "no disponible"}


def _load_workbook(path: Path) -> list[tuple[str, list[list[object]]]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        return [(ws.title, [list(row) for row in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    if suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(path)
        out = []
        for sheet in book.sheets():
            rows = []
            for r in range(sheet.nrows):
                rows.append([sheet.cell_value(r, c) if sheet.cell_type(r, c) != xlrd.XL_CELL_EMPTY else None for c in range(sheet.ncols)])
            out.append((sheet.name, rows))
        return out
    raise SchemaDriftError(f"unsupported workbook extension: {suffix}")


def _cell_identity(sheet: str, row: int, col: int) -> str:
    return f"{sheet}!R{row + 1}C{col + 1}"


def _is_missing(value: object) -> bool:
    if value is None:
        return False
    return norm(value) in _MISSING_MARKERS or str(value).strip().lower() in _MISSING_MARKERS


def _publication_date_from_url(url: str) -> str | None:
    # Companion files conventionally encode MM_YY. This is publication month,
    # not a fabricated day; keep day unknown in the normalized product.
    m = re.search(r"_(\d{2})_(\d{2})(?:\D|$)", url)
    if not m:
        return None
    month, yy = int(m.group(1)), int(m.group(2))
    if 1 <= month <= 12:
        return f"20{yy:02d}-{month:02d}"
    return None


def _context_indicator(rows: list[list[object]], row: int, sheet: str, lookback: int = 10):
    # Resolve the nearest explicit anchor. Looking across the entire window can
    # accidentally merge adjacent human-formatted table blocks.
    for rr in range(row, max(-1, row - lookback - 1), -1):
        found = {ind for value in rows[rr] if (ind := match_indicator(value))}
        if len(found) == 1:
            return next(iter(found))
        if len(found) > 1:
            # Historical matrix headers legitimately place several indicator
            # labels on one row. That row cannot identify a single block;
            # continue toward a more specific anchor (or let the caller's
            # no-observations/schema-drift gate fail closed).
            continue
    sheet_ind = match_indicator(sheet)
    return sheet_ind


def _context_period(rows: list[list[object]], row: int, styles: tuple[str, ...], sheet: str, lookback: int = 10):
    for rr in range(row, max(-1, row - lookback - 1), -1):
        found = {p for value in rows[rr] if (p := canonical_period(value, allowed_styles=styles))}
        if len(found) == 1:
            return next(iter(found))
        if len(found) > 1:
            raise SchemaDriftError(f"ambiguous period anchor in {sheet} row {rr + 1}: {sorted(found)}")
    return canonical_period(sheet, allowed_styles=styles)


def _context_geography(rows: list[list[object]], row: int, sheet: str, lookback: int = 10):
    for rr in range(row, max(-1, row - lookback - 1), -1):
        found = {geo for value in rows[rr] if (geo := match_geography(value))}
        if len(found) == 1:
            return next(iter(found))
        if len(found) > 1:
            return None
    return match_geography(sheet)


def _wide_period_columns(rows: list[list[object]], row: int) -> dict[int, str]:
    """Resolve explicit stacked year/quarter headers above a wide matrix row."""
    year_row = None
    quarter_row = None
    for rr in range(row - 1, max(-1, row - 20), -1):
        values = [norm(v) for v in rows[rr]]
        if year_row is None and any(re.fullmatch(r"ano ?\d{4}", v) for v in values):
            year_row = rr
        if quarter_row is None and any(re.search(r"[1-4].*trimestre", v) for v in values):
            quarter_row = rr
        if year_row is not None and quarter_row is not None:
            break
    if year_row is None or quarter_row is None:
        return {}
    periods: dict[int, str] = {}
    current_year: str | None = None
    for c, value in enumerate(rows[year_row]):
        match = re.fullmatch(r"ano ?(\d{4})", norm(value))
        if match:
            current_year = match.group(1)
        quarter = re.search(r"([1-4]).*trimestre", norm(rows[quarter_row][c]))
        if current_year and quarter:
            periods[c] = f"{current_year}-Q{quarter.group(1)}"
    return periods


def _make_observation(period: str, geo, indicator, raw: object, snapshot: SourceSnapshot, sheet: str, row: int, col: int) -> Observation | None:
    value = decimal_or_none(raw)
    if value is None and not _is_missing(raw):
        return None
    return Observation(
        period=period,
        geography_level=geo.geography_level,
        geography_id=geo.geography_id,
        geography_label=geo.geography_label,
        indicator_id=indicator.indicator_id,
        indicator_label=indicator.indicator_label,
        value=value,
        unit=indicator.unit,
        numerator_universe=indicator.numerator_universe,
        denominator_universe=indicator.denominator_universe,
        source_id=snapshot.source_id,
        source_publication_date=_publication_date_from_url(snapshot.url),
        source_snapshot_sha256=snapshot.sha256,
        source_cell_identity=_cell_identity(sheet, row, col),
        value_status="observed" if value is not None else "missing_source",
        source_geography_code=geo.source_code,
        source_geography_label=geo.geography_label,
    )


def parse_rate_workbook(path: Path, snapshot: SourceSnapshot) -> list[Observation]:
    if snapshot.layout_id not in ("indec_historical_matrix_v1", "indec_current_matrix_v2"):
        raise SchemaDriftError(f"unsupported rates layout_id: {snapshot.layout_id}")
    styles = _LAYOUT_PERIOD_STYLES[snapshot.layout_id]
    observations: list[Observation] = []
    for sheet, rows in _load_workbook(path):
        # Official reports repeat derived summaries in several tables.  The
        # governed rate cells come from the principal total table plus its
        # current-period geographic table; other tables can use different
        # rounding/universe conventions and are not interchangeable sources.
        if snapshot.layout_id == "indec_current_matrix_v2" and sheet.lower().startswith("cuadro"):
            allowed = {"cuadro 1", "cuadro 1.1", "cuadro 1.2", "cuadro 2", "cuadro 3.1", "cuadro 3.2"}
            if sheet.lower() not in allowed:
                continue
        if not rows:
            continue
        max_cols = max((len(r) for r in rows), default=0)
        rows = [r + [None] * (max_cols - len(r)) for r in rows]
        consumed_cells: set[tuple[int, int]] = set()

        # Orientation A: geographies across columns, periods down rows, one indicator block.
        for r, row in enumerate(rows):
            geo_cols = [(c, match_geography(v)) for c, v in enumerate(row)]
            geo_cols = [(c, g) for c, g in geo_cols if g]
            if not geo_cols:
                continue
            indicator = _context_indicator(rows, r, sheet)
            if not indicator:
                continue
            # A single total-31 header is acceptable; other single geographies are too
            # ambiguous for block detection.
            if len(geo_cols) == 1 and geo_cols[0][1].geography_id != "total_31_agglomerates":
                continue
            for rr in range(r + 1, min(len(rows), r + 120)):
                if rr > r + 1 and any(match_indicator(v) for v in rows[rr]):
                    break
                period = next((canonical_period(v, allowed_styles=styles) for v in rows[rr][:8] if canonical_period(v, allowed_styles=styles)), None)
                if not period:
                    continue
                for c, geo in geo_cols:
                    if c >= len(rows[rr]):
                        continue
                    obs = _make_observation(period, geo, indicator, rows[rr][c], snapshot, sheet, rr, c)
                    if obs:
                        observations.append(obs); consumed_cells.add((rr, c))

        # Orientation B: periods across columns, geographies down rows, one indicator block.
        for r, row in enumerate(rows):
            period_cols = [(c, canonical_period(v, allowed_styles=styles)) for c, v in enumerate(row)]
            period_cols = [(c, p) for c, p in period_cols if p]
            if len(period_cols) < 2:
                continue
            indicator = _context_indicator(rows, r, sheet)
            if not indicator:
                continue
            for rr in range(r + 1, min(len(rows), r + 80)):
                if rr > r + 1 and any(match_indicator(v) for v in rows[rr]):
                    break
                geo = next((match_geography(v) for v in rows[rr][:8] if match_geography(v)), None)
                if not geo:
                    continue
                for c, period in period_cols:
                    if (rr, c) in consumed_cells:
                        continue
                    obs = _make_observation(period, geo, indicator, rows[rr][c], snapshot, sheet, rr, c)
                    if obs:
                        observations.append(obs); consumed_cells.add((rr, c))

        # Orientation C: indicators across columns, geographies down rows, period in block title.
        for r, row in enumerate(rows):
            ind_cols = [(c, match_indicator(v)) for c, v in enumerate(row)]
            ind_cols = [(c, i) for c, i in ind_cols if i]
            if len(ind_cols) < 2:
                continue
            period = _context_period(rows, r, styles, sheet)
            if not period:
                continue
            for rr in range(r + 1, min(len(rows), r + 80)):
                if rr > r + 1 and any(match_indicator(v) for v in rows[rr]):
                    break
                geo = next((match_geography(v) for v in rows[rr][:8] if match_geography(v)), None)
                if not geo:
                    continue
                for c, indicator in ind_cols:
                    if (rr, c) in consumed_cells:
                        continue
                    obs = _make_observation(period, geo, indicator, rows[rr][c], snapshot, sheet, rr, c)
                    if obs:
                        observations.append(obs); consumed_cells.add((rr, c))

        # Orientation D0: indicators down rows, explicit year/quarter columns
        # above them, and geography supplied by the table title/header. This
        # is the wide quarterly matrix used by the current INDEC workbook.
        for r, row in enumerate(rows):
            indicators = [indicator for value in row if (indicator := match_indicator(value))]
            if len(indicators) != 1:
                continue
            period_cols = _wide_period_columns(rows, r)
            if len(period_cols) < 2:
                continue
            geo = _context_geography(rows, r, sheet)
            if not geo:
                continue
            indicator = indicators[0]
            for c, period in period_cols.items():
                if c >= len(row) or (r, c) in consumed_cells:
                    continue
                obs = _make_observation(period, geo, indicator, row[c], snapshot, sheet, r, c)
                if obs:
                    observations.append(obs); consumed_cells.add((r, c))

        # Orientation D: indicators across columns, periods down rows, geography in block title.
        for r, row in enumerate(rows):
            ind_cols = [(c, match_indicator(v)) for c, v in enumerate(row)]
            ind_cols = [(c, i) for c, i in ind_cols if i]
            if len(ind_cols) < 2:
                continue
            geo = _context_geography(rows, r, sheet)
            if not geo:
                continue
            for rr in range(r + 1, min(len(rows), r + 120)):
                if rr > r + 1 and any(match_indicator(v) for v in rows[rr]):
                    break
                period = next((canonical_period(v, allowed_styles=styles) for v in rows[rr][:8] if canonical_period(v, allowed_styles=styles)), None)
                if not period:
                    continue
                for c, indicator in ind_cols:
                    if (rr, c) in consumed_cells:
                        continue
                    obs = _make_observation(period, geo, indicator, rows[rr][c], snapshot, sheet, rr, c)
                    if obs:
                        observations.append(obs); consumed_cells.add((rr, c))

    # L1's governed product is the four principal indicators.  The same
    # workbook also contains component/pressure tables whose labels can map to
    # non-core indicators; retain them in the source workbook, but do not let
    # them collide with the principal-rate cells.
    observations = [o for o in observations if o.indicator_id in CORE_INDICATORS]

    # The current publication repeats the same official matrix in two tables
    # (Cuadro 3.1 and 3.2).  Treat identical repeats as one observation while
    # retaining both source-cell identities; conflicting repeats fail closed.
    by_key: dict[tuple[str, str, str, str], Observation] = {}
    for obs in observations:
        prior = by_key.get(obs.key)
        if prior is None:
            by_key[obs.key] = obs
            continue
        if (prior.value, prior.value_status) != (obs.value, obs.value_status):
            # Regional/current summary tables publish rounded total-31 values;
            # the principal total table is the governed source for that cell.
            # Preserve the alternate cell identity without replacing the
            # higher-precision principal observation.
            if obs.geography_id == "total_31_agglomerates" and (
                "Cuadro 1.1!" in prior.source_cell_identity
                or "Cuadro 1!" in prior.source_cell_identity
            ):
                prior.source_cell_identity = ";".join(
                    dict.fromkeys(prior.source_cell_identity.split(";") + [obs.source_cell_identity])
                )
                continue
            raise SchemaDriftError(
                f"conflicting repeated official cell {obs.key}: "
                f"{prior.value!r} ({prior.source_cell_identity}) vs "
                f"{obs.value!r} ({obs.source_cell_identity})"
            )
        prior.source_cell_identity = ";".join(
            dict.fromkeys(prior.source_cell_identity.split(";") + [obs.source_cell_identity])
        )
    observations = list(by_key.values())
    if not observations:
        raise SchemaDriftError(
            f"no recognized official labor cells in {path.name}; layout={snapshot.layout_id}. "
            "This is a schema-drift failure, not permission to guess column positions."
        )
    return observations


_METRIC_ALIASES = {
    "cv": ("coeficiente de variacion", "cv"),
    "ci90_low": ("limite inferior 90", "limite inferior del 90", "inferior 90", "li"),
    "ci90_high": ("limite superior 90", "limite superior del 90", "superior 90", "ls"),
}


def _quality_metric(value: object) -> str | None:
    key = norm(value)
    for metric, aliases in _METRIC_ALIASES.items():
        if any(norm(alias) == key or (len(norm(alias)) >= 8 and norm(alias) in key) for alias in aliases):
            return metric
    return None


def parse_quality_workbook(path: Path, snapshot: SourceSnapshot) -> list[QualityRecord]:
    if snapshot.layout_id != "indec_quality_v1":
        raise SchemaDriftError(f"unsupported quality layout_id: {snapshot.layout_id}")
    styles = _LAYOUT_PERIOD_STYLES[snapshot.layout_id]
    records: list[QualityRecord] = []
    for sheet, rows in _load_workbook(path):
        # The quality workbook also publishes component tables (for example
        # demanding/non-demanding subemployment).  Those are not the core
        # indicator cells joined to the rate release; their abbreviated sheet
        # titles otherwise resolve to the generic subemployment indicator.
        sheet_key = norm(sheet)
        if "demandante" in sheet_key or "no demand" in sheet_key:
            continue
        if not rows:
            continue
        max_cols = max((len(r) for r in rows), default=0)
        rows = [r + [None] * (max_cols - len(r)) for r in rows]
        for r, row in enumerate(rows):
            metric_cols = [(c, _quality_metric(v)) for c, v in enumerate(row)]
            metric_cols = [(c, m) for c, m in metric_cols if m]
            data_start = r + 1
            # Current INDEC quality workbooks split the metric header across
            # two rows: CV/"Intervalo de confianza" followed by LI/LS. Join
            # only that explicit adjacent header row; do not infer columns
            # from arbitrary table values.
            if {m for _, m in metric_cols} == {"cv"} and r + 1 < len(rows):
                next_metrics = [
                    (c, _quality_metric(v)) for c, v in enumerate(rows[r + 1])
                ]
                next_metrics = [(c, m) for c, m in next_metrics if m]
                metric_cols.extend(next_metrics)
                data_start = r + 2
            # The continuation row (LI/LS) is itself encountered on the
            # next loop iteration.  It belongs to the preceding CV header;
            # parsing it again would duplicate every quality cell.
            if metric_cols and {m for _, m in metric_cols} <= {"ci90_low", "ci90_high"}:
                continue
            if len({m for _, m in metric_cols}) < 2:
                continue
            indicator = _context_indicator(rows, r, sheet)
            period = _context_period(rows, r, styles, sheet)
            if not indicator or not period:
                continue
            for rr in range(data_start, min(len(rows), r + 80)):
                geo = next((match_geography(v) for v in rows[rr][:8] if match_geography(v)), None)
                if not geo:
                    continue
                values = {metric: decimal_or_none(rows[rr][c]) for c, metric in metric_cols}
                if not any(v is not None for v in values.values()):
                    continue
                cells = ",".join(_cell_identity(sheet, rr, c) for c, _ in metric_cols)
                records.append(QualityRecord(
                    period=period,
                    geography_id=geo.geography_id,
                    indicator_id=indicator.indicator_id,
                    cv=values.get("cv"),
                    ci90_low=values.get("ci90_low"),
                    ci90_high=values.get("ci90_high"),
                    source_id=snapshot.source_id,
                    source_snapshot_sha256=snapshot.sha256,
                    source_cell_identity=cells,
                ))
    if not records:
        raise SchemaDriftError(
            f"no recognized CV/CI cells in {path.name}; layout={snapshot.layout_id}. "
            "Quality metadata is never synthesized when the source schema is not recognized."
        )
    return records
